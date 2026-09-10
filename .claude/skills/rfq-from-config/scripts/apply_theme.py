#!/usr/bin/env python3
"""Apply design tokens to a generated RFQ document.

Every visual decision reads from assets/theme.json, so adopting a different
design system means editing that file and not this one.

Two things are deliberately left alone:

* **The letterhead.** Page headers and footers carry the company's logo
  artwork. Restyling type over the top of fixed images is how a letterhead
  stops lining up, and the logos already define the brand on the page.
* **Data shading.** The spare-parts appendix marks rows with colour to mean
  something. Recolouring those cells to match a palette would erase the
  meaning while looking tidier.

So this restyles typography, section headings and table header rows - the
document's structure - and leaves the parts that carry information alone.
"""
from __future__ import annotations

import json
import os
import re
import sys

try:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import RGBColor
except ImportError:
    sys.exit("python-docx is required:  pip install python-docx")

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_THEME = os.path.join(HERE, "..", "assets", "theme.json")

# First rows worth treating as a header, matched on the words they contain.
# Anything unrecognised is left as it is: the letterhead blocks and the terms
# table open with content rather than column titles, and shading those would
# turn a value into a heading.
HEADER_SIGNATURES = (
    {"module", "description"},        # Configuration
    {"part number", "description"},   # Pricing
    {"service type", "description"},  # Service contract options
    {"group", "description"},         # Spare parts
)


def load_theme(path: str = DEFAULT_THEME) -> dict:
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.upper())


def shade(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def distinct_cells(row) -> list:
    out = []
    for cell in row.cells:
        if not out or cell._tc is not out[-1]._tc:
            out.append(cell)
    return out


def _is_header_row(row) -> bool:
    words = {c.text.strip().lower() for c in distinct_cells(row)}
    return any(sig <= words for sig in HEADER_SIGNATURES)


def find_header_row(table):
    """The row carrying column titles, which is not always the first.

    The spare-parts appendix opens with a merged title row and puts its column
    names underneath, so looking only at row 0 leaves that table unthemed.
    """
    for row in table.rows[:3]:
        if _is_header_row(row):
            return row
    return None


def find_config_table(doc):
    """The Configuration table - the only one whose subheadings we write."""
    for table in doc.tables:
        if table.rows and _is_header_row(table.rows[0]):
            words = {c.text.strip().lower() for c in distinct_cells(table.rows[0])}
            if "module" in words:
                return table
    return None


def _looks_like_heading(paragraph) -> bool:
    """A standalone section title rather than a sentence of body text.

    Headings in this template are short, bold and unpunctuated - 'Configuration',
    'Pricing', 'Terms & conditions'. Requiring all three keeps the closing
    paragraphs and the bulleted highlights out of it.
    """
    text = paragraph.text.strip()
    if not text or len(text) > 60 or text.endswith((".", ",", ":")):
        return False
    runs = [r for r in paragraph.runs if r.text.strip()]
    return bool(runs) and all(r.bold for r in runs)


def apply_theme(doc, theme: dict) -> dict:
    """Restyle the document in place. Returns a summary of what changed."""
    font = theme.get("font", {})
    colors = theme.get("color", {})
    table_tokens = theme.get("table", {})
    body_font = font.get("body")
    heading_font = font.get("heading", body_font)
    heading_color = colors.get("heading_text")

    counts = {"runs_refonted": 0, "headings": 0, "header_rows": 0, "subheadings": 0}

    # --- typography ---------------------------------------------------------
    # Only the family is set. Sizes carry the document's hierarchy - the
    # subject line is larger than the body on purpose - and flattening them to
    # one size would destroy that while claiming to be a restyle.
    if body_font:
        for paragraph in _body_paragraphs(doc):
            for run in paragraph.runs:
                run.font.name = body_font
                rPr = run._r.get_or_add_rPr()
                rFonts = rPr.find(qn("w:rFonts"))
                if rFonts is None:
                    rFonts = OxmlElement("w:rFonts")
                    rPr.insert(0, rFonts)
                for attr in ("w:ascii", "w:hAnsi", "w:cs"):
                    rFonts.set(qn(attr), body_font)
                counts["runs_refonted"] += 1

    # --- section headings ---------------------------------------------------
    for paragraph in doc.paragraphs:
        if not _looks_like_heading(paragraph):
            continue
        for run in paragraph.runs:
            if heading_color:
                run.font.color.rgb = _rgb(heading_color)
            if heading_font:
                run.font.name = heading_font
        counts["headings"] += 1

    # --- table header rows and class subheadings ----------------------------
    header_fill = table_tokens.get("header_fill")
    header_text = table_tokens.get("header_text")
    sub_fill = table_tokens.get("subheading_fill")
    sub_text = table_tokens.get("subheading_text")

    for table in doc.tables:
        if not table.rows:
            continue
        header_row = find_header_row(table) if header_fill else None
        if header_row is not None:
            for cell in distinct_cells(header_row):
                shade(cell, header_fill)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        if header_text:
                            run.font.color.rgb = _rgb(header_text)
            counts["header_rows"] += 1

    # Class subheadings are only ever written into the Configuration table.
    # Treating every merged row as one would also catch the pricing table's
    # TOTAL line and the spare-parts title - rows that mean something else and
    # already carry their own emphasis.
    config = find_config_table(doc) if sub_fill else None
    if config is not None:
        for row in config.rows[1:]:
            cells = distinct_cells(row)
            if len(cells) == 1 and cells[0].text.strip():
                shade(cells[0], sub_fill)
                for paragraph in cells[0].paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        if sub_text:
                            run.font.color.rgb = _rgb(sub_text)
                counts["subheadings"] += 1

    return counts


def _body_paragraphs(doc):
    """Every paragraph in the document body - headers and footers excluded."""
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            seen = set()
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                yield from cell.paragraphs


def main(argv=None) -> int:
    import argparse
    from docx import Document

    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("document", help="the .docx to restyle in place")
    ap.add_argument("-t", "--theme", default=DEFAULT_THEME)
    args = ap.parse_args(argv)

    doc = Document(args.document)
    theme = load_theme(args.theme)
    counts = apply_theme(doc, theme)
    doc.save(args.document)

    print(f"Themed {args.document} using {os.path.basename(args.theme)}")
    print(f"  source          : {theme.get('source', 'unspecified')}")
    print(f"  headings        : {counts['headings']}")
    print(f"  table headers   : {counts['header_rows']}")
    print(f"  class subheads  : {counts['subheadings']}")
    print(f"  runs re-fonted  : {counts['runs_refonted']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
