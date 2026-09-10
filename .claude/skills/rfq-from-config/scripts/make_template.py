#!/usr/bin/env python3
"""Re-derive the blank RFQ template from a filled example offer.

When the company letterhead changes, the fastest safe route to a new template
is to take a real offer that already uses the new letterhead and strip the
customer-specific content out of it. Rebuilding the Word file from scratch
would mean recreating logos, headers, footers, styles and the spare-parts
appendix by hand.

This blanks the letterhead fields into single-run {{TOKEN}} placeholders and
reduces the Configuration table to its header plus one empty prototype row,
which is what build_rfq.py clones for each module.

Usage:
    python3 make_template.py FILLED_OFFER.docx -o ../assets/rfq_template.docx
"""
from __future__ import annotations

import argparse
import copy
import os
import sys

try:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
except ImportError:
    sys.exit("python-docx is required:  pip install python-docx")

# Label -> token. The label text is matched case-insensitively against the
# start of a paragraph, so a reworded letterhead still lines up.
FIELDS = [
    ("to:",              "To:  ",              "CUSTOMER"),
    ("protocol number:", "Protocol number:  ", "PROTOCOL"),
    ("date",             "Date:  ",            "DATE"),
    ("subject:",         "Subject:  ",         "SUBJECT"),
    ("description:",     "Description:     ",  "DESCRIPTION_1"),
]


def set_para(para, parts: list[str]) -> None:
    """Rewrite a paragraph as one run per entry in `parts`.

    Every run inherits the first existing run's formatting. Keeping each token
    in its own run is what lets build_rfq.py swap it without disturbing the
    surrounding font - Word otherwise splits text across runs unpredictably.
    """
    runs = para.runs
    if not runs:
        return
    proto = runs[0]._r
    for run in runs[1:]:
        run._r.getparent().remove(run._r)
    first = para.runs[0]
    first.text = parts[0]
    tail = first._r
    for extra in parts[1:]:
        new = copy.deepcopy(proto)
        for node in new.findall(qn("w:t")):
            new.remove(node)
        node = OxmlElement("w:t")
        node.set(qn("xml:space"), "preserve")
        node.text = extra
        new.append(node)
        tail.addnext(new)
        tail = new


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("source", help="a filled offer using the desired letterhead")
    ap.add_argument("-o", "--output", default="rfq_template.docx")
    args = ap.parse_args(argv)

    doc = Document(args.source)
    done = set()

    # --- letterhead fields -------------------------------------------------
    for table in doc.tables[:6]:
        for row in table.rows:
            for cell in row.cells:
                for idx, para in enumerate(cell.paragraphs):
                    text = para.text.strip().lower()
                    for label, prefix, token in FIELDS:
                        if token in done or not text.startswith(label):
                            continue
                        set_para(para, [prefix, "{{%s}}" % token])
                        done.add(token)
                        # The machine description runs onto a second line.
                        if token == "DESCRIPTION_1" and idx + 1 < len(cell.paragraphs):
                            set_para(cell.paragraphs[idx + 1], ["{{DESCRIPTION_2}}"])
                            done.add("DESCRIPTION_2")

    # --- configuration table ----------------------------------------------
    config = None
    for table in doc.tables:
        header = [c.text.strip().lower() for c in table.rows[0].cells]
        if "module" in header and "description" in header:
            config = table
            break
    if config is None:
        sys.exit("No Configuration table found (need a header row with "
                 "'Module' and 'Description').")

    for row in list(config.rows)[2:]:
        row._tr.getparent().remove(row._tr)
    for cell in config.rows[1].cells:
        set_para(cell.paragraphs[0], [""])
        for extra in cell.paragraphs[1:]:
            extra._p.getparent().remove(extra._p)

    # A 60-row table crosses pages; repeating the header keeps it readable.
    trPr = config.rows[0]._tr.get_or_add_trPr()
    if trPr.find(qn("w:tblHeader")) is None:
        trPr.append(OxmlElement("w:tblHeader"))

    doc.save(args.output)
    print(f"Wrote {os.path.abspath(args.output)}")
    print("  placeholders: " + (", ".join("{{%s}}" % t for t in sorted(done)) or "none"))
    missing = {t for _, _, t in FIELDS} - done
    if missing:
        print("  NOT found in the source (fill these by hand): "
              + ", ".join("{{%s}}" % t for t in sorted(missing)))
    print("  configuration table reduced to header + 1 prototype row")
    return 0


if __name__ == "__main__":
    sys.exit(main())
